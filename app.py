import io
import csv
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import json
import requests
import streamlit as st


OLLAMA_URL = "http://localhost:11434/api/chat"
MODEL = "qwen2.5:3b"

st.set_page_config(
    page_title="StudyMate AI",
    page_icon="📚",
    layout="wide"
)

# ---------- Modern Frontend Styling ----------

st.markdown("""
<style>

/* Main background */
.stApp {
    background: linear-gradient(135deg, #f8faff, #f5f3ff);
}

/* Main content width */
.block-container {
    max-width: 1200px;
    padding-top: 2rem;
    padding-bottom: 3rem;
}

/* Font and typography */
html, body, [class*="st-"] {
    font-family: "Segoe UI", Arial, sans-serif;
}

/* Headings */
h1, h2, h3 {
    color: #172554 !important;
    font-weight: 700 !important;
}

/* Main title */
h1 {
    text-align: center;
    font-size: 2.7rem !important;
    letter-spacing: -1px;
}

/* Paragraphs */
.stMarkdown p {
    color: #475569;
    line-height: 1.7;
}

/* Buttons */
.stButton > button,
.stDownloadButton > button {
    background: linear-gradient(135deg, #4f46e5, #6366f1);
    color: white;
    border: none;
    border-radius: 12px;
    padding: 0.65rem 1.2rem;
    font-weight: 600;
    transition: all 0.2s ease;
}

.stButton > button:hover,
.stDownloadButton > button:hover {
    background: #4338ca;
    color: white;
    transform: translateY(-2px);
    box-shadow: 0 5px 15px #4f46e533;
}

/* Notes and text fields */
.stTextArea textarea,
.stTextInput input {
    background-color: white;
    border: 1px solid #dbe4f0;
    border-radius: 12px;
}

/* File uploader */
[data-testid="stFileUploader"] {
    background: white;
    border: 1px dashed #a5b4fc;
    border-radius: 16px;
    padding: 1rem;
}

/* Tabs */
.stTabs [data-baseweb="tab-list"] {
    gap: 10px;
    border-bottom: 1px solid #e2e8f0;
}

.stTabs [data-baseweb="tab"] {
    border-radius: 10px 10px 0 0;
    padding: 12px 16px;
    font-weight: 600;
}

/* Alerts */
[data-testid="stAlert"] {
    border-radius: 12px;
}

/* Metrics */
[data-testid="stMetric"] {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 14px;
    padding: 16px;
}

/* Mobile layout */
@media (max-width: 768px) {
    .block-container {
        padding: 1rem;
    }

    h1 {
        font-size: 2rem !important;
    }
}

</style>
""", unsafe_allow_html=True)

# ---------- App Header ----------

st.title("📚 StudyMate AI")

st.markdown("""
<p style="
    text-align:center;
    color:#64748b;
    font-size:1.1rem;
    margin-top:-10px;
">
Turn your study notes into understanding.
Learn smarter with your private, local AI assistant.
</p>
""", unsafe_allow_html=True)

st.markdown("---")

# ---------- Local AI helper ----------

def ask_ai(prompt):
    response = requests.post(
        OLLAMA_URL,
        json={
            "model": MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "You are a helpful study assistant. "
                        "Give accurate, simple, student-friendly answers."
                    )
                },
                {"role": "user", "content": prompt}
            ],
            "stream": False
        },
        timeout=300
    )
    response.raise_for_status()
    return response.json()["message"]["content"]


def extract_json(text):
    # Handle JSON wrapped in Markdown code fences
    text = text.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[-1]
        if text.endswith("```"):
            text = text[:-3]
    return json.loads(text.strip())


# ---------- Session state ----------

if "quiz" not in st.session_state:
    st.session_state.quiz = None

if "quiz_result" not in st.session_state:
    st.session_state.quiz_result = None


# ---------- Document Upload ----------

def extract_text(uploaded_file):
    file_name = uploaded_file.name.lower()
    text = ""

    if file_name.endswith(".pdf"):
        pdf = PdfReader(uploaded_file)
        text = "\n".join(
            page.extract_text() or ""
            for page in pdf.pages
        )

    elif file_name.endswith(".docx"):
        doc = Document(uploaded_file)
        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    elif file_name.endswith(".pptx"):
        presentation = Presentation(uploaded_file)
        slides_text = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    slides_text.append(shape.text)

        text = "\n".join(slides_text)

    elif file_name.endswith(".xlsx"):
        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True
        )
        rows = []

        for sheet in workbook.worksheets:
            rows.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) for value in row
                    if value is not None
                ]
                if values:
                    rows.append(" | ".join(values))

        text = "\n".join(rows)
        workbook.close()

    elif file_name.endswith(".csv"):
        content = uploaded_file.getvalue().decode(
            "utf-8-sig", errors="replace"
        )
        text = content

    elif file_name.endswith((".txt", ".md")):
        text = uploaded_file.getvalue().decode(
            "utf-8-sig", errors="replace"
        )

    return text



# ---------- Study Workspace Dashboard ----------

# Initialize session state BEFORE using it
if "study_notes" not in st.session_state:
    st.session_state["study_notes"] = ""

if "loaded_document" not in st.session_state:
    st.session_state["loaded_document"] = None

if "quiz" not in st.session_state:
    st.session_state["quiz"] = None

if "quiz_result" not in st.session_state:
    st.session_state["quiz_result"] = None

left_col, right_col = st.columns([1.7, 1], gap="large")

with left_col:
    st.subheader("📂 Your Study Workspace")

    uploaded_file = st.file_uploader(
        "Upload your study material",
        type=["pdf", "docx", "pptx", "xlsx", "csv", "txt", "md"]
    )

    if uploaded_file is not None:
        file_signature = (uploaded_file.name, uploaded_file.size)

        if st.session_state["loaded_document"] != file_signature:
            try:
                uploaded_text = extract_text(uploaded_file)

                if uploaded_text.strip():
                    st.session_state["study_notes"] = uploaded_text
                    st.session_state["loaded_document"] = file_signature
                    st.session_state["quiz"] = None
                    st.session_state["quiz_result"] = None
                    st.session_state.pop("performance_analysis", None)

                    st.success(f"Uploaded: {uploaded_file.name}")
                else:
                    st.warning("No readable text found in the document.")

            except Exception as e:
                st.error(f"Could not read document: {e}")

    notes = st.text_area(
        "📖 Study Notes",
        height=260,
        placeholder="Upload a document or paste notes here...",
        key="study_notes"
    )

with right_col:
    st.subheader("✨ Your Study Tools")

    st.markdown("""
    **📝 Quiz & Performance**

    Test your knowledge with five multiple-choice questions.

    ---

    **📚 Summarize Notes**

    Turn lengthy notes into clear summaries.

    ---

    **🎯 Topic Q&A**

    Prepare for exams with important questions and answers.
    """)

    st.info(
        "💡 Add your notes on the left, then choose a study tool below."
    )

# ---------- Create Tabs ----------

tab_quiz, tab_summary, tab_topic = st.tabs(
    [
        "📝 Quiz & Performance",
        "📚 Summarize Notes",
        "🎯 Topic Q&A"
    ]
)


# ---------- Quiz section ----------

with tab_quiz:
    st.subheader("Take a Quiz")
    st.write("Create a five-question multiple-choice quiz from your notes.")

    if st.button("🎯 Generate Quiz", key="generate_quiz"):
        if not notes.strip():
            st.warning("Please paste your study notes first.")
        else:
            prompt = f"""
Create exactly 5 multiple-choice questions based ONLY on these notes.

Return valid JSON only, with no Markdown or extra text, in this format:
{{
  "questions": [
    {{
      "question": "Question text",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "answer": 0,
      "explanation": "Why the correct answer is right"
    }}
  ]
}}

Rules:
- Each question must have exactly 4 options.
- answer must be an integer from 0 to 3.
- Exactly one option must be correct.
- Questions must be answerable from the notes.
- Do not reveal the correct answer in the question.
- Make questions clear and suitable for a student.

NOTES:
{notes}
"""
            try:
                with st.spinner("AI is creating your quiz..."):
                    raw = ask_ai(prompt)
                    quiz = extract_json(raw)

                questions = quiz.get("questions", [])

                if len(questions) != 5:
                    st.error(
                        "AI did not create exactly 5 questions. "
                        "Please try generating the quiz again."
                    )
                elif not all(
                    isinstance(q.get("question"), str)
                    and isinstance(q.get("options"), list)
                    and len(q["options"]) == 4
                    and isinstance(q.get("answer"), int)
                    and not isinstance(q.get("answer"), bool)
                    and 0 <= q["answer"] < 4
                    and all(isinstance(option, str) for option in q["options"])
                    for q in questions
                ):
                    st.error(
                        "The generated quiz format was invalid. "
                        "Please try again."
                    )
                else:
                    st.session_state.quiz = questions
                    st.session_state.quiz_result = None
                    st.rerun()

            except (requests.RequestException, KeyError, ValueError) as e:
                st.error(f"Could not generate quiz: {e}")

    quiz = st.session_state.quiz

    if quiz:
        st.divider()
        st.write(f"**Questions: {len(quiz)}**")

        if st.session_state.quiz_result is None:
            with st.form("quiz_form"):
                answers = {}

                for i, q in enumerate(quiz):
                    st.markdown(f"**Q{i + 1}. {q['question']}**")

                    answers[i] = st.radio(
                        f"Choose your answer for question {i + 1}",
                        options=range(4),
                        format_func=lambda x, opts=q["options"]: opts[x],
                        index=None,
                        key=f"answer_{i}",
                        label_visibility="collapsed"
                    )
                    st.write("")

                submitted = st.form_submit_button(
                    "✅ Submit Quiz",
                    use_container_width=True
                )

            if submitted:
                unanswered = [
                    i + 1 for i, answer in answers.items()
                    if answer is None
                ]

                if unanswered:
                    st.warning(
                        "Please answer all questions. Unanswered: "
                        + ", ".join(map(str, unanswered))
                    )
                else:
                    score = sum(
                        answers[i] == q["answer"]
                        for i, q in enumerate(quiz)
                    )

                    st.session_state.quiz_result = {
                        "answers": answers,
                        "score": score
                    }
                    st.rerun()

        # ---------- Results and AI analysis ----------

        result = st.session_state.quiz_result

        if result is not None:
            score = result["score"]
            answers = result["answers"]
            percentage = round(score / len(quiz) * 100)

            st.subheader("📊 Your Results")

            col1, col2, col3 = st.columns(3)
            col1.metric("Score", f"{score}/{len(quiz)}")
            col2.metric("Percentage", f"{percentage}%")
            col3.metric("Incorrect", len(quiz) - score)

            if percentage >= 80:
                st.success("Great work! You understand much of this material.")
            elif percentage >= 50:
                st.info("Good effort! Review the questions you missed.")
            else:
                st.warning("Keep practising. Review your notes and try again.")

            st.subheader("📝 Answer Review")

            for i, q in enumerate(quiz):
                chosen = answers[i]
                correct = q["answer"]

                if chosen == correct:
                    st.markdown(f"**Q{i + 1}: ✅ Correct**")
                else:
                    st.markdown(f"**Q{i + 1}: ❌ Incorrect**")

                st.write(q["question"])
                st.write(f"Your answer: {q['options'][chosen]}")
                st.write(f"Correct answer: {q['options'][correct]}")
                st.caption(q.get("explanation", ""))
                st.divider()

            if st.button("🧠 Analyse My Performance", use_container_width=True):
                review = []

                for i, q in enumerate(quiz):
                    review.append({
                        "question": q["question"],
                        "student_answer": q["options"][answers[i]],
                        "correct_answer": q["options"][q["answer"]],
                        "was_correct": answers[i] == q["answer"]
                    })

                prompt = f"""
Analyse this student's quiz performance using the data below.

Score: {score}/{len(quiz)}
Percentage: {percentage}%

Question-by-question results:
{json.dumps(review, indent=2)}

Provide:
1. Overall performance in simple language.
2. What the student appears to understand well.
3. Mistakes and the concepts to review.
4. Three specific study recommendations.
5. A short revision plan.

Use evidence from the actual answers. Do not claim a broad weakness
from just one question. If the evidence is limited, say so.
Do not change the score or invent additional answers.
"""
                try:
                    with st.spinner("AI is analysing your answers..."):
                        analysis = ask_ai(prompt)

                    st.session_state.performance_analysis = analysis

                except requests.RequestException as e:
                    st.error(f"Could not analyse performance: {e}")

            if st.session_state.get("performance_analysis"):
                st.subheader("🎓 AI Performance Analysis")
                st.markdown(st.session_state.performance_analysis)

            if st.button("🔄 Try a New Quiz"):
                st.session_state.quiz = None
                st.session_state.quiz_result = None
                st.session_state.pop("performance_analysis", None)

                for i in range(5):
                    st.session_state.pop(f"answer_{i}", None)

                st.rerun()

# ---------- Summary section ----------

with tab_summary:
    st.subheader("Summarize Your Notes")

    if st.button("📚 Summarize", key="summarize"):
        if not notes.strip():
            st.warning("Please paste your notes first.")
        else:
            try:
                with st.spinner("Summarizing your notes..."):
                    summary = ask_ai(
                        "Summarize these notes in simple language using "
                        "clear bullet points. Stay faithful to the notes.\n\n"
                        + notes
                    )

                st.markdown(summary)

            except requests.RequestException as e:
                st.error(f"Could not summarize notes: {e}")

# ---------- Topic-Based Important Questions ----------

with tab_topic:
    st.subheader("🎯 Important Questions & Answers")

    st.write(
        "Enter any topic and get 10 important questions "
        "with answers for exam preparation."
    )

    topic = st.text_input(
        "Enter your topic",
        placeholder="e.g. Java OOP, DBMS, Photosynthesis",
        key="important_topic"
    )

    level = st.selectbox(
        "Select difficulty level",
        ["Beginner", "Intermediate", "Advanced"],
        key="important_level"
    )

    if st.button("Generate 10 Questions", key="topic_qa"):

        if not topic.strip():
            st.warning("Please enter a topic first.")

        else:
            prompt = f"""
You are an educational assistant helping a student prepare for exams.

Topic: {topic}
Difficulty: {level}

Return exactly 10 different important questions about the topic.

IMPORTANT RULES:
1. Every item MUST contain both a complete question and its answer.
2. Do not return answers alone.
3. Questions must be specific to the given topic.
4. Use clear, student-friendly language.
5. Include examples or code where useful.
6. Do not repeat questions.
7. Return valid JSON only, without Markdown code fences.
8. Use exactly this JSON structure:

{{
  "questions": [
    {{
      "question": "Write the complete question here",
      "answer": "Write the complete answer here"
    }}
  ]
}}

The questions array must contain exactly 10 objects.
Each object must contain both question and answer.
"""

            try:
                with st.spinner(
                    "AI is preparing your questions and answers..."
                ):
                    result = ask_ai(prompt)

                # Convert the AI response into structured data
                data = extract_json(result)
                questions = data.get("questions", [])

                # Check whether the AI returned complete questions
                if (
                    len(questions) != 10
                    or any(
                        not isinstance(item, dict)
                        or not str(item.get("question", "")).strip()
                        or not str(item.get("answer", "")).strip()
                        for item in questions
                    )
                ):
                    st.error(
                        "The AI did not return all 10 complete "
                        "questions and answers. Please try again."
                    )

                else:
                    st.success(
                        f"10 questions generated for: {topic}"
                    )

                    # Build readable output for display and download
                    output = []

                    for i, item in enumerate(questions, start=1):
                        question = item["question"].strip()
                        answer = item["answer"].strip()

                        st.markdown(f"### Q{i}. {question}")
                        st.markdown(f"**Answer:** {answer}")
                        st.divider()

                        output.append(
                            f"Q{i}. {question}\nAnswer: {answer}"
                        )

                    final_text = "\n\n".join(output)

                    st.download_button(
                        label="📥 Download Questions & Answers",
                        data=final_text,
                        file_name="important_questions.txt",
                        mime="text/plain"
                    )

            except json.JSONDecodeError:
                st.error(
                    "The AI returned an invalid format. "
                    "Please click Generate 10 Questions again."
                )

            except requests.exceptions.ConnectionError:
                st.error(
                    "Cannot connect to Ollama. "
                    "Make sure Ollama is running."
                )

            except requests.exceptions.RequestException as e:
                st.error(f"Could not generate questions: {e}")
